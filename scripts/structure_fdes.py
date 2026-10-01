import pandas as pd
from openpyxl import load_workbook

fdes_path = "frameworks/raw/fdes-essat-2026.xlsx"
wb = load_workbook(fdes_path)

def split_also_in(text):
    """Separate an '(also in ...)' note from a text column.
    Returns two columns: the text with the note removed, and the
    referenced FDES code(s), with multiple codes separated by '; '.
    """
    # Capture everything after "also in" up to the closing bracket
    codes = text.str.extract(r"also in\s*([^)]+)")[0]
    # "2.5.1.b and 2.5.3.b" -> "2.5.1.b; 2.5.3.b"
    codes = codes.str.replace(" and ", "; ")
    # Remove the note (and the space before it) from the original text
    cleaned = text.str.replace(r"\s*\(also in [^)]*\)", "", regex=True)
    return cleaned, codes

def process_component(sheet_name):
    """Read one FDES Basic Set component sheet and return one row per statistic."""
    
    comp = pd.read_excel(fdes_path, sheet_name=sheet_name, header=None)
    ws = wb[sheet_name]

    excel_rows = []
    tiers = []

    for row in ws.iter_rows():
        group = row[0]
        statistic = row[1]

        if statistic.value is not None:
            cell = statistic
        else:
            cell = group

        if cell.font.bold:
            tier = 1
        elif cell.font.italic:
            tier = 3
        else:
            tier = 2

        excel_rows.append(cell.row)
        tiers.append(tier)
        #print(cell.row, cell.value, tier)

    tier_df = pd.DataFrame({
        "excel_row": excel_rows, 
        "tier": tiers
    })

    comp = comp.iloc[:,0:3]
    comp["excel_row"] = comp.index + 1

    comp = comp.merge(tier_df, on="excel_row", how="left")
    comp = comp.rename(columns={0: "col_a", 1: "col_b", 2: "unit"})

    # Break out columns
    comp["component"] = comp["col_a"].where(
        comp["col_a"].str.startswith("Component", na=False)
    )
    comp["component"] = comp["component"].ffill()

    comp["subcomponent"] = comp["col_a"].where(
        comp["col_a"].str.startswith("Sub-component", na=False)
    )
    comp["subcomponent"] = comp["subcomponent"].ffill()

    comp["topic"] = comp["col_a"].where(
        comp["col_a"].str.startswith("Topic", na=False)
    )
    comp["topic"] = comp["topic"].ffill()

    is_header = comp["col_a"].str.startswith(("Component", "Sub-component", "Topic"), na=False)
    comp["statistic_group"] = comp["col_a"].where(~is_header)
    comp["statistic_group"] = comp["statistic_group"].ffill()
    comp["statistic"] = comp["col_b"]

    # Keep rows with statistics
    has_statistic = comp["statistic"].str.match(r"^\d+\.", na=False)
    is_standalone = comp["col_a"].str.match(r"^[a-z]\.", na=False) & comp["statistic"].isna()
    comp = comp[has_statistic | is_standalone]

    # Detail codes
    comp["component_code"] = comp["component"].str.extract(r"Component ([\d.]+):")
    comp["subcomponent_code"] = comp["subcomponent"].str.extract(r"Sub-component ([\d.]+):")
    comp["topic_code"] = comp["topic"].str.extract(r"Topic ([\d.]+):")

    comp["statistic_group_letter"] = comp["statistic_group"].str.extract(r"^([a-z])\.")
    comp["statistic_group_code"] = comp["topic_code"] + "." + comp["statistic_group_letter"]

    comp["statistic_number"] = comp["statistic"].str.extract(r"^(\d+)\.")
    comp["statistic_code"] = (
        comp["statistic_group_code"] + "." + comp["statistic_number"]
    ).fillna(comp["statistic_group_code"])

    comp["fdes_code"] = comp["statistic_code"]
    comp = comp.drop(columns=["statistic_group_letter", "statistic_number"])

    # Detail names
    comp["component_name"] = comp["component"].str.extract(r":\s*(.+)")
    comp["subcomponent_name"] = comp["subcomponent"].str.extract(r":\s*(.+)")
    comp["topic_name"] = comp["topic"].str.extract(r":\s*(.+)")

    # Extract notes for statistic/groups also found in other areas
    comp["statistic_group"], group_ref = split_also_in(comp["statistic_group"])
    comp["statistic"], stat_ref = split_also_in(comp["statistic"])
    comp["also_in"] = stat_ref.fillna(group_ref)

    comp["statistic_group_name"] = comp["statistic_group"].str.extract(r"^[a-z]\.\s*(.+)")
    comp["statistic_name"] = comp["statistic"].str.extract(r"^\d+\.\s*(.+)")

    # Clean columns
    comp = comp.drop(columns = ["col_a", "col_b", "excel_row"])

    # Trim leading/trailing spaces in every text column
    for col in comp.columns:
        if col != "tier":
            comp[col] = comp[col].str.strip()

    comp["statistic_group_name"] = comp["statistic_group_name"].str.rstrip(":") # Remove trailing : of some statistic groups

    # Create compiled statistic name
    comp["fdes_statistic"] = (
        comp["statistic_group_name"] + ": " + comp["statistic_name"]
    ).fillna(comp["statistic_group_name"])

    comp = comp[[
        "fdes_code", "fdes_statistic",
        "component_code", "component_name",
        "subcomponent_code", "subcomponent_name",
        "topic_code", "topic_name",
        "statistic_group_code", "statistic_group_name",
        "statistic_code", "statistic_name",
        "tier",
        "unit",
        "also_in",
        "component", "subcomponent", "topic", "statistic_group", "statistic",
    ]]

    #comp.to_csv("frameworks/structured/comp_check.csv", index=False, encoding="utf-8-sig")
    return comp


sheets = [f"Component {i}" for i in range(1, 7)]
fdes = pd.concat([process_component(s) for s in sheets], ignore_index=True)

print(len(fdes))                    # expect 458
print(fdes["fdes_code"].is_unique)  # expect True
print(fdes.isna().sum())

#fdes.to_csv("frameworks/structured/fdes_table.csv", index=False, encoding="utf-8-sig")